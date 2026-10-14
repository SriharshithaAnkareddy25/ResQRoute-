"""Build the ResQRoute research paper (IJACSA-style, two-column) -> paper/ResQRoute_paper.docx

All numbers come from outputs/ produced by scripts 00-09. Citations are written as {key} and numbered
in order of first appearance.

  python paper/build_paper.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "outputs" / "figures"
OUT = ROOT / "paper" / "ResQRoute_paper.docx"

REFS = {
    "dong2013": 'L. Dong and J. Shan, "A comprehensive review of earthquake-induced building damage detection with remote sensing techniques," ISPRS J. Photogramm. Remote Sens., vol. 84, pp. 85-99, 2013.',
    "xbd": 'R. Gupta et al., "xBD: A dataset for assessing building damage from satellite imagery," arXiv:1911.09296, 2019.',
    "xbdcvpr": 'R. Gupta et al., "Creating xBD: A dataset for assessing building damage from satellite imagery," in Proc. IEEE/CVF CVPR Workshops, 2019.',
    "zhou2019": 'Z. Zhou, X. Chen, E. Li, L. Zeng, K. Luo, and J. Zhang, "Edge intelligence: Paving the last mile of artificial intelligence with edge computing," Proc. IEEE, vol. 107, no. 8, pp. 1738-1762, 2019.',
    "viola": 'P. Viola and M. Jones, "Rapid object detection using a boosted cascade of simple features," in Proc. IEEE CVPR, 2001.',
    "idk": 'X. Wang, Y. Luo, D. Crankshaw, A. Tumanov, F. Yu, and J. E. Gonzalez, "IDK cascades: Fast deep learning by learning not to overthink," in Proc. UAI, 2018.',
    "bolukbasi": 'T. Bolukbasi, J. Wang, O. Dekel, and V. Saligrama, "Adaptive neural networks for efficient inference," in Proc. ICML, 2017.',
    "xu2019": 'J. Z. Xu, W. Lu, Z. Li, P. Khaitan, and V. Zaytseva, "Building damage detection in satellite imagery using convolutional neural networks," arXiv:1910.06444, 2019.',
    "weber": 'E. Weber and H. Kane, "Building disaster damage assessment in satellite imagery with multi-temporal fusion," arXiv:2004.05525, 2020.',
    "rescuenet": 'R. Gupta and M. Shah, "RescueNet: Joint building segmentation and damage assessment from satellite imagery," in Proc. ICPR, 2020.',
    "bdanet": 'Y. Shen et al., "BDANet: Multiscale convolutional neural network with cross-directional attention for building damage assessment from satellite images," IEEE Trans. Geosci. Remote Sens., vol. 60, 2022.',
    "mobilenetv3": 'A. Howard et al., "Searching for MobileNetV3," in Proc. IEEE/CVF ICCV, 2019.',
    "resnet": 'K. He, X. Zhang, S. Ren, and J. Sun, "Deep residual learning for image recognition," in Proc. IEEE CVPR, 2016.',
    "efficientnet": 'M. Tan and Q. V. Le, "EfficientNet: Rethinking model scaling for convolutional neural networks," in Proc. ICML, 2019.',
    "imagenet": 'J. Deng, W. Dong, R. Socher, L.-J. Li, K. Li, and L. Fei-Fei, "ImageNet: A large-scale hierarchical image database," in Proc. IEEE CVPR, 2009.',
    "branchynet": 'S. Teerapittayanon, B. McDanel, and H. T. Kung, "BranchyNet: Fast inference via early exiting from deep neural networks," in Proc. ICPR, 2016.',
    "msdnet": 'G. Huang, D. Chen, T. Li, F. Wu, L. van der Maaten, and K. Q. Weinberger, "Multi-scale dense networks for resource efficient image classification," in Proc. ICLR, 2018.',
    "sdn": 'Y. Kaya, S. Hong, and T. Dumitras, "Shallow-deep networks: Understanding and mitigating network overthinking," in Proc. ICML, 2019.',
    "skipnet": 'X. Wang, F. Yu, Z.-Y. Dou, T. Darrell, and J. E. Gonzalez, "SkipNet: Learning dynamic routing in convolutional networks," in Proc. ECCV, 2018.',
    "dynsurvey": 'Y. Han, G. Huang, S. Song, L. Yang, H. Wang, and Y. Wang, "Dynamic neural networks: A survey," IEEE Trans. Pattern Anal. Mach. Intell., vol. 44, no. 11, pp. 7436-7456, 2022.',
    "noscope": 'D. Kang, J. Emmons, F. Abuzaid, P. Bailis, and M. Zaharia, "NoScope: Optimizing neural network queries over video at scale," Proc. VLDB Endow., vol. 10, no. 11, 2017.',
    "neurosurgeon": 'Y. Kang et al., "Neurosurgeon: Collaborative intelligence between the cloud and mobile edge," in Proc. ASPLOS, 2017.',
    "edgent": 'E. Li, Z. Zhou, and X. Chen, "Edge intelligence: On-demand deep learning model co-inference with device-edge synergy," in Proc. ACM SIGCOMM Workshop on Mobile Edge Communications, 2018.',
    "frugalml": 'L. Chen, M. Zaharia, and J. Zou, "FrugalML: How to use ML prediction APIs more accurately and cheaply," in Proc. NeurIPS, 2020.',
    "frugalgpt": 'L. Chen, M. Zaharia, and J. Zou, "FrugalGPT: How to use large language models while reducing cost and improving performance," arXiv:2305.05176, 2023.',
    "geifman": 'Y. Geifman and R. El-Yaniv, "Selective classification for deep neural networks," in Proc. NeurIPS, 2017.',
    "madras": 'D. Madras, T. Pitassi, and R. Zemel, "Predict responsibly: Improving fairness and accuracy by learning to defer," in Proc. NeurIPS, 2018.',
    "mozannar": 'H. Mozannar and D. Sontag, "Consistent estimators for learning to defer to an expert," in Proc. ICML, 2020.',
    "guo": 'C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, "On calibration of modern neural networks," in Proc. ICML, 2017.',
    "hendrycks": 'D. Hendrycks and K. Gimpel, "A baseline for detecting misclassified and out-of-distribution examples in neural networks," in Proc. ICLR, 2017.',
    "elkan": 'C. Elkan, "The foundations of cost-sensitive learning," in Proc. IJCAI, 2001.',
    "szegedy": 'C. Szegedy, V. Vanhoucke, S. Ioffe, J. Shlens, and Z. Wojna, "Rethinking the inception architecture for computer vision," in Proc. IEEE CVPR, 2016.',
    "muller": 'R. Muller, S. Kornblith, and G. Hinton, "When does label smoothing help?," in Proc. NeurIPS, 2019.',
    "friedman": 'J. H. Friedman, "Greedy function approximation: A gradient boosting machine," Ann. Statist., vol. 29, no. 5, pp. 1189-1232, 2001.',
    "lightgbm": 'G. Ke et al., "LightGBM: A highly efficient gradient boosting decision tree," in Proc. NeurIPS, 2017.',
    "sklearn": 'F. Pedregosa et al., "Scikit-learn: Machine learning in Python," J. Mach. Learn. Res., vol. 12, pp. 2825-2830, 2011.',
    "pytorch": 'A. Paszke et al., "PyTorch: An imperative style, high-performance deep learning library," in Proc. NeurIPS, 2019.',
    "efron": 'B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap. New York, NY, USA: Chapman & Hall, 1993.',
}
_order: list[str] = []


def cite(text: str) -> str:
    def rep(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        nums = []
        for k in keys:
            if k not in REFS:
                raise KeyError(k)
            if k not in _order:
                _order.append(k)
            nums.append(_order.index(k) + 1)
        return "[" + "], [".join(str(n) for n in sorted(nums)) + "]"
    return re.sub(r"\{([a-z][a-z0-9_]*(?:, ?[a-z][a-z0-9_]*)*)\}", rep, text)


# ---------------------------------------------------------------- formatting helpers
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(3)
for s in doc.sections:
    s.page_height, s.page_width = Inches(11.69), Inches(8.27)  # A4
    s.left_margin = s.right_margin = Inches(0.7)
    s.top_margin = s.bottom_margin = Inches(0.75)


def set_cols(section, n, space=0.25):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space * 1440)))


def para(text="", bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, before=0, indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = Inches(0.15)
    if text:
        add_runs(p, cite(text), bold, italic, size)
    return p


def add_runs(p, text, bold=False, italic=False, size=None):
    # **bold** and *italic* inline markup
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        r = p.add_run(part.strip("*") if part.startswith("*") else part)
        r.bold = bold or part.startswith("**")
        r.italic = italic or (part.startswith("*") and not part.startswith("**"))
        if size:
            r.font.size = Pt(size)


SEC = [0]
SUB = [0]


def heading(text):
    SEC[0] += 1
    SUB[0] = 0
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"][SEC[0] - 1]
    p = para(f"{roman}. {text.upper()}", align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, size=10)
    p.paragraph_format.keep_with_next = True
    p.runs[0].font.small_caps = True


def sub(text):
    SUB[0] += 1
    p = para(f"{chr(64 + SUB[0])}. {text}", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2)
    p.paragraph_format.keep_with_next = True


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(1)
        add_runs(p, cite(it))


def numbered(items):
    for i, it in enumerate(items, 1):
        p = para(indent=False)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        add_runs(p, f"{i})  " + cite(it))


def equation(text, num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.italic = True
    p.add_run(f"\t({num})")


TAB = [0]


def table(caption, header, rows, widths, font=7.5):
    TAB[0] += 1
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"][TAB[0] - 1]
    c = para(f"TABLE {roman}.  {caption.upper()}", align=WD_ALIGN_PARAGRAPH.CENTER, size=8, before=6, after=2)
    c.runs[0].font.small_caps = True
    c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        add_runs(cell.paragraphs[0], h, bold=True, size=font)
        shade(cell, "E8E8E8")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            add_runs(cell.paragraphs[0], str(v), size=font)
    for row in t.rows:
        for c_ in row.cells:
            for p_ in c_.paragraphs:
                p_.paragraph_format.keep_with_next = True
        for j, w in enumerate(widths):
            row.cells[j].width = Inches(w)
            for p in row.cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Inches(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
    para(after=2)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


FIGN = [0]


def figure(path, caption, width=3.25):
    FIGN[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    para(f"Fig. {FIGN[0]}. {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, size=8, after=6)


# ---------------------------------------------------------------- title block (one column)
t = para("ResQRoute: A Risk- and Resource-Aware Adaptive Inference Controller for Post-Disaster Building Damage Assessment",
         bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
para("[Author 1]¹, [Author 2]², [Author 3]³", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
para("Department of Computer Science and Engineering, [Institution Name], [City], [Country]",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

# ---------------------------------------------------------------- body (two columns)
body = doc.add_section(WD_SECTION.CONTINUOUS)
set_cols(body, 2)

p = para(indent=False)
add_runs(p, "**Abstract**—", size=9)
add_runs(p, cite(
    "Rapid assessment of building damage from satellite imagery is critical after floods, hurricanes and wildfires, "
    "but running a large deep network on every building is slow, energy-hungry and often impossible on field devices "
    "with intermittent connectivity. Conventional model cascades escalate an input to a larger model whenever the small "
    "model's confidence falls below a fixed threshold, ignoring both the device's current resources and the very "
    "different consequences of different errors. This paper proposes ResQRoute, a learned routing controller that, for "
    "each building, chooses among five actions: accept a lightweight prediction, escalate to one of two larger models, "
    "defer to a priority queue, or flag for human review. The decision jointly uses prediction uncertainty, cheap "
    "image-quality cues, the device's latency budget, memory, network and battery state, and an asymmetric penalty "
    "that weights missed severe damage most heavily. Three ImageNet-pretrained classifiers (MobileNetV3-Small, "
    "ResNet-18 and EfficientNet-B0) were fine-tuned on 6,000 balanced pre/post building crops from three disasters "
    "in the xBD dataset, and their predictions, CPU latency and memory were logged for every crop. On a tile-disjoint "
    "test set evaluated under 6,816 simulated device states, the controller reached a macro-F1 of 0.631 at an average "
    "of 27 ms per image, matching the 0.610 of always running EfficientNet-B0 (126 ms) or ResNet-18 (235 ms) at "
    "roughly one-fifth to one-ninth of the computation, while exceeding the device budget in only 0.5% of decisions "
    "versus 33% for a fixed-confidence cascade and reducing dangerous severe-to-no-damage errors from 7.4% to 4.6%."), size=9)
p = para(indent=False, after=8)
add_runs(p, "**Keywords**—", size=9)
add_runs(p, "*Adaptive inference; model cascade; building damage assessment; xBD; edge computing; cost-sensitive learning; learning to defer; disaster response*", size=9)

# I. Introduction
heading("Introduction")
para("After a major natural disaster, emergency managers must decide within hours where to send search-and-rescue "
     "teams, engineers and supplies. Very-high-resolution satellite imagery makes it possible to assess thousands of "
     "buildings remotely {dong2013}, and public benchmarks such as xBD {xbd, xbdcvpr} have enabled convolutional "
     "neural networks to grade building damage automatically. In practice, however, these assessments are often run on "
     "laptops or edge devices close to the affected area, where computation, memory, battery and network bandwidth are "
     "scarce and change over time {zhou2019}.")
para("Not every building is equally hard to assess. A clearly intact house in a sharp image can be classified "
     "correctly by a lightweight network, whereas a partially collapsed roof under smoke or shadow may require a "
     "stronger model or a human analyst. Model cascades exploit this by running a cheap model first and escalating "
     "only uncertain inputs {viola, idk, bolukbasi}. Most cascades, though, use a single fixed confidence threshold. "
     "They do not consider whether the device can currently afford the larger model, and they treat all errors "
     "equally, although predicting *no damage* for a destroyed building is far more dangerous than confusing *minor* "
     "with *no damage*.")
para("This paper addresses the question: **which model should process each image under the current resource and "
     "safety constraints?** The contributions are:")
numbered([
    "**A five-action routing controller** (accept, escalate to either of two larger models, defer, human review) "
    "trained as a supervised classifier on oracle routes derived from the logged behaviour of three base models.",
    "**Joint decision features** combining prediction uncertainty, image-quality cues, and the device state "
    "(latency budget, free memory, network, battery, queue length), together with an **asymmetric safety weighting** "
    "that penalises missed severe damage most heavily.",
    "**A deferred-escalation queue** ordered by a priority score P = U × R × D × A, so that postponed, uncertain and "
    "potentially dangerous buildings are processed first when resources return.",
    "**A reproducible, laptop-scale evaluation** on xBD with tile-disjoint splits, real CPU latency and memory "
    "measurements, five baselines, and bootstrap confidence intervals.",
])
para("Section II surveys related work, Section III describes the proposed method, Section IV the experimental setup, "
     "Section V the results, Section VI the limitations, and Section VII concludes.")

# II. Literature review
heading("Literature Survey")
sub("Building damage assessment from remote sensing")
para("Early damage mapping relied on change detection between pre- and post-event imagery and on manual photo-"
     "interpretation; Dong and Shan {dong2013} review these techniques for earthquakes. The xBD dataset {xbd, xbdcvpr} "
     "provides about 850,000 building polygons over 19 events with a four-level damage scale (no, minor, major "
     "damage, destroyed) and was released with the xView2 challenge. Deep-learning approaches on xBD include siamese "
     "and multi-temporal fusion CNNs {xu2019, weber}, joint segmentation and damage classification (RescueNet) "
     "{rescuenet}, and attention-based multiscale networks such as BDANet {bdanet}. These works aim to maximise "
     "accuracy with a single, fixed model; none of them adapts the amount of computation per building or to the "
     "deployment device.")
sub("Efficient convolutional networks")
para("Mobile-oriented architectures reduce inference cost: MobileNetV3 {mobilenetv3} combines depthwise-separable "
     "convolutions with architecture search, ResNet {resnet} introduced residual learning and remains a common "
     "medium-sized backbone, and EfficientNet {efficientnet} scales depth, width and resolution jointly. All three "
     "are typically initialised from ImageNet {imagenet}. Choosing one of them fixes the accuracy-cost trade-off "
     "for every input; ResQRoute instead uses all three as interchangeable experts.")
sub("Adaptive and conditional inference")
para("Cascades date back to the Viola-Jones detector {viola}. In deep learning, early-exit networks such as BranchyNet "
     "{branchynet}, MSDNet {msdnet} and Shallow-Deep Networks {sdn} attach classifiers to intermediate layers and stop "
     "when confident; SkipNet {skipnet} learns to skip layers; a survey of such dynamic networks is given in {dynsurvey}. "
     "Model-level selection is closest to this work: Bolukbasi *et al.* {bolukbasi} learn a policy that decides whether "
     "to evaluate a more expensive network, IDK cascades {idk} learn when a model should say \"I don't know\", and "
     "NoScope {noscope} cascades cheap specialised models in front of an expensive detector for video. FrugalML "
     "{frugalml} and FrugalGPT {frugalgpt} learn which paid prediction API to call under a monetary budget. "
     "Neurosurgeon {neurosurgeon} and Edgent {edgent} adapt by partitioning a single network between device and cloud "
     "according to network conditions. These methods optimise an accuracy-cost trade-off but treat all errors as "
     "equally costly and, except for the edge-partitioning works, do not condition on a changing device state.")
sub("Selective prediction, deferral and calibration")
para("Selective classification lets a model abstain on inputs where it is likely to be wrong {geifman}; learning-to-"
     "defer frameworks train a classifier jointly with a rejector that hands inputs to a human expert {madras, mozannar}. "
     "Such decisions rely on well-calibrated confidence, which modern networks often lack {guo}; label smoothing "
     "{szegedy, muller} is a simple way to soften over-confident outputs, and maximum softmax probability is a common "
     "baseline for detecting misclassified or out-of-distribution inputs {hendrycks}. Cost-sensitive learning "
     "{elkan} formalises unequal misclassification costs. These lines of work handle abstention or unequal costs, "
     "but not jointly with multi-model routing under resource constraints.")
sub("Research gap")
para("Table I summarises the survey. To the best of our knowledge, no prior work on satellite damage assessment "
     "combines (i) learned selection among several models, (ii) conditioning on the current device resources, "
     "(iii) asymmetric safety costs, and (iv) deferral and human-review actions in a single controller. ResQRoute "
     "targets exactly this combination.")
Y, N, P_ = "Yes", "No", "Partly"
table("Comparison of related approaches",
      ["Approach", "Learned model choice", "Device resources", "Asym. risk", "Defer / human"],
      [["xBD CNNs {xu2019, weber, rescuenet, bdanet}".replace("{xu2019, weber, rescuenet, bdanet}", cite("{xu2019, weber, rescuenet, bdanet}")), N, N, N, N],
       [cite("Fixed-threshold cascade {viola}"), N, N, N, N],
       [cite("Early-exit nets {branchynet, msdnet, sdn}"), P_, N, N, N],
       [cite("Adaptive nets / IDK {bolukbasi, idk}"), Y, N, N, P_],
       [cite("NoScope, FrugalML {noscope, frugalml}"), Y, P_ + " (budget)", N, N],
       [cite("Neurosurgeon, Edgent {neurosurgeon, edgent}"), N, Y, N, N],
       [cite("Learning to defer {madras, mozannar}"), N, N, P_, Y],
       ["**ResQRoute (ours)**", Y, Y, Y, Y]],
      [1.35, 0.55, 0.55, 0.45, 0.5])

# III. Method
heading("Proposed Methodology")
sub("System overview")
para("ResQRoute contains three damage classifiers of increasing cost and a routing controller. The pipeline is "
     "small model → controller → action: for each building crop the small model always runs "
     "first. The controller then observes the small model's output, image-quality features and the device state, and "
     "selects one of five actions a ∈ {accept, escalate-EfficientNet, escalate-ResNet-18, defer, human-review}. "
     "Escalation targets are ordered by their *measured* CPU cost rather than by nominal size (Section V).")
sub("Data preparation")
para("The labelled training release of xBD {xbd} (2,799 pre/post image pairs of 1024 × 1024 pixels from 10 events) "
     "was indexed into 162,787 building polygons. Because the minor-damage class is rare in most events, three events "
     "were selected that together cover all damage levels: Hurricane Michael and Hurricane Matthew (wind damage of all "
     "severities, most minor-damage examples) and the Santa Rosa wildfire (a second hazard type with many destroyed "
     "buildings). Hurricane Harvey was excluded because its many major-damage labels stem from flooding that is "
     "largely invisible from above. Major damage and destroyed were merged into *severe damage*, giving three classes; "
     "unclassified buildings and polygons smaller than 150 px² were removed.")
para("From 43,148 usable buildings, 2,000 per class were sampled (severe split equally between major and destroyed), "
     "yielding 6,000 buildings from 668 tiles. For each, the polygon's bounding box was enlarged by 20% per side "
     "(minimum 32 px), cropped from both the pre- and post-event image, and letterboxed to 224 × 224. The data were "
     "split 70/15/15 (4,299 / 849 / 852) with stratification by class **and grouping by source tile**, so that "
     "neighbouring buildings from the same scene never appear in different splits.")
sub("Base classifiers")
para("MobileNetV3-Small {mobilenetv3}, ResNet-18 {resnet} and EfficientNet-B0 {efficientnet}, pretrained on ImageNet "
     "{imagenet}, were fine-tuned on the post-event crop with PyTorch {pytorch}: AdamW (learning rate 3×10⁻⁴), cosine "
     "schedule, flips, 90° rotations and colour jitter, class-weighted cross-entropy with label smoothing 0.1 "
     "{szegedy, muller} to temper over-confident probabilities that the controller consumes, and checkpoint selection "
     "by validation macro-F1 (early stopping after 6 epochs without improvement, maximum 20).")
sub("Controller features")
para("At decision time the controller sees only information that is actually available before escalation:")
bullets([
    "**Uncertainty:** the small model's class probabilities, maximum confidence, top-2 margin and predictive entropy "
    "H = −Σ pₖ log pₖ.",
    "**Image quality:** variance of the Laplacian (blur), mean and standard deviation of grey level (brightness, "
    "contrast), fraction of the crop that is image rather than padding (building size), mean absolute pre/post "
    "difference (change magnitude), and the shares of near-black and near-white pixels (shadow, cloud or smoke proxies).",
    "**Device state:** latency budget, free memory, network availability, battery level, queue length, and whether "
    "each larger model fits the current budget and memory.",
])
sub("Oracle routes and asymmetric safety")
para("Training labels are derived from the logged behaviour of the three models on each crop i under device state s. "
     "Let cₘ be whether model m is correct and fₘ(s) whether it fits the budget and memory. The oracle route is:")
equation("a*(i,s) = accept if c_small; else the cheapest m with cₘ ∧ fₘ(s); else defer if ∃m: cₘ; else human-review", 1)
para("with human review replaced by deferral when the network is unavailable. To encode the asymmetric cost of errors "
     "{elkan}, each training example is weighted by w = 1 for non-severe buildings, w = 3 for severe buildings, and "
     "w = 5 for severe buildings that the small model misclassifies. The controller is a histogram-based gradient-"
     "boosted tree ensemble {friedman, lightgbm} from scikit-learn {sklearn}. At inference two hard guards apply: a "
     "model that does not fit the device is never chosen (the request is deferred instead), and human review is "
     "converted to deferral when offline.")
sub("Deferred escalation queue")
para("Deferred buildings are re-processed by EfficientNet-B0 when resources return, in order of the priority")
equation("Pᵢ = Uᵢ × Rᵢ × Dᵢ × Aᵢ", 2)
para("where Uᵢ is the normalised entropy (uncertainty), Rᵢ the small model's probability of severe damage (risk of a "
     "dangerous miss), Dᵢ = 0.5·p_minor + p_severe the expected severity, and Aᵢ = 0.5 + 0.5·(building size fraction) "
     "an area-importance term.")

# IV. Experimental setup
heading("Experimental Setup")
para("Training ran on Kaggle's free CPU sessions (4 cores). All inference measurements were made on a laptop CPU with "
     "one thread, batch size 1, after 10 warm-up passes; each crop's latency is the median of three timed forward "
     "passes, and memory is the peak resident-set increase for loading a model and running one forward pass, measured "
     "in a fresh process. The controller was trained on the 849 validation crops (on which the base models were not "
     "trained, so their outputs are honest), each paired with 8 random device states, and evaluated on the 852 test "
     "crops × 8 states = 6,816 decisions. Device states draw the latency budget from {60, 180, 400} ms, free memory "
     "from {150, 250, 500} MB, network availability with probability 0.7, and battery uniformly in [0.05, 1]; below "
     "20% battery the budget is halved.")
para("Baselines: always running each single model; a **fixed-confidence cascade** (accept when confidence ≥ t, else "
     "escalate, ignoring resources); and a **resource-aware cascade without risk** (accept when confident, else the "
     "cheapest model that fits, else defer). Both thresholds were tuned on validation (t = 0.74 and t = 0.52). Metrics: "
     "accuracy, macro-F1, severe-damage recall, dangerous false-negative rate (truly severe predicted as no-damage), "
     "average latency and peak memory, the share of decisions exceeding the device budget, and the distribution of "
     "actions. Human-reviewed buildings are counted as correctly labelled; results without this credit are also "
     "reported. Confidence intervals use a crop-level bootstrap {efron} with 500 resamples.")

# V. Results
heading("Results and Discussion")
sub("Base classifiers")
table("Base classifiers on the test split (852 crops)",
      ["Model", "Acc.", "Macro-F1", "Severe recall", "CPU ms", "Mem. MB"],
      [["MobileNetV3-Small", "0.604", "0.597", "0.398", "22.9", "98.5"],
       ["ResNet-18", "0.622", "0.610", "0.377", "219.8", "172.1"],
       ["EfficientNet-B0", "0.617", "0.610", "0.419", "105.3", "121.5"]],
      [1.15, 0.4, 0.5, 0.5, 0.45, 0.45])
para("Table II shows that the three models are close in accuracy (validation macro-F1 0.776, 0.795 and 0.778), "
     "while their costs differ by an order of magnitude. On a CPU, ResNet-18 is the slowest and most memory-hungry "
     "model despite not being the most accurate, so cost cannot be inferred from nominal model size. The dominant "
     "error for all models is confusing severe with minor damage. At least one of the three models is correct on "
     "75.7% of test crops versus 62.2% for the best single model, which indicates substantial room for routing.")
sub("Routing performance")
table("Routing policies on the test split (6,816 decisions)",
      ["Policy", "Macro-F1", "Dang. FN %", "Avg ms", "Over budget %"],
      [["Always MobileNetV3-S", "0.597", "6.3", "23", "0.5"],
       ["Always EfficientNet-B0", "0.610", "4.2", "126", "37.0"],
       ["Always ResNet-18", "0.610", "9.5", "235", "78.3"],
       ["Fixed-confidence cascade", "0.612", "7.4", "113", "32.5"],
       ["Resource-aware, no risk", "0.605", "4.6", "33", "0.4"],
       ["**ResQRoute (ours)**", "**0.631**", "**4.6**", "**27**", "**0.5**"]],
      [1.3, 0.5, 0.5, 0.45, 0.6])
figure(FIG / "cost_vs_quality.png", "Macro-F1 versus average CPU time per image on the test split.")
para("Table III and Fig. 1 show the central result. The controller attains the highest macro-F1 (0.631) at an "
     "average of 27 ms per image, about one-fifth of always running EfficientNet-B0 and one-ninth of ResNet-18. It "
     "respects the device in 99.5% of decisions, whereas the fixed-confidence cascade exceeds the budget or memory in "
     "32.5% of them, because it escalates regardless of resources. Its dangerous false-negative rate (4.6%) is lower "
     "than that of the small model (6.3%) and the fixed cascade (7.4%), and close to the best single model (4.2%). "
     "The controller accepted the small model in 82.8% of decisions, escalated in 4.0%, deferred 9.1% and sent 4.0% "
     "to human review.")
figure(FIG / "dangerous_fn.png", "Share of truly severe buildings predicted as no-damage (lower is safer).")
sub("Adaptation to resources")
figure(FIG / "actions_by_budget.png", "Controller actions by latency budget.")
para("Fig. 3 shows the controller's behaviour changing with the device: under a tight 60 ms budget it almost never "
     "escalates (0.5%) and defers 12.9% of buildings; under a relaxed 400 ms budget it escalates 6.1% and defers only "
     "5.7%. Accuracy rises accordingly from 0.623 to 0.652 while average latency stays between 23 and 31 ms.")
sub("Deferred queue and learned features")
para("Of the 619 deferred decisions, 43.7% of those in the top half of the priority queue were truly severe, against "
     "34.5% in the bottom half, so equation (2) brings dangerous buildings forward. Permutation importance on the "
     "validation set ranks network availability, the small model's probability of severe damage, the top-2 margin, "
     "whether each larger model fits the budget, and the dark-pixel share and brightness of the crop as the most "
     "influential features, confirming that the controller combines uncertainty, resources and image quality.")
sub("Statistical robustness")
para("A crop-level bootstrap gives a 95% confidence interval of [−0.003, +0.044] for the macro-F1 difference between "
     "ResQRoute and always running EfficientNet-B0, and [−0.006, +0.043] against the fixed-confidence cascade. Without "
     "crediting human review, ResQRoute's macro-F1 is 0.607 versus 0.610 for EfficientNet-B0 (difference −0.003, 95% "
     "CI [−0.026, +0.019]). The defensible conclusion is therefore that ResQRoute **matches** the accuracy of the large "
     "models at a fraction of their cost while staying within the device budget and reducing dangerous errors "
     "relative to a conventional cascade; the accuracy gain itself is mostly due to human review and is not "
     "statistically significant on this test set.")

# VI. Limitations
heading("Limitations")
bullets([
    "**Simulated device states.** Model latencies and memory are measured, but budgets, network and battery are "
    "sampled from fixed distributions; field traces would strengthen the evaluation.",
    "**Small controller training set.** Only 849 held-out crops are available for the controller; out-of-fold "
    "predictions over the training split would provide about five times more.",
    "**Split difficulty.** The test split contains half of its severe buildings from Hurricane Matthew, whose imagery "
    "is the blurriest; test macro-F1 (≈0.61) is therefore well below validation (≈0.78).",
    "**Post-image only.** The base classifiers use only the post-event crop; xView2 solutions show that comparing "
    "pre- and post-event images substantially improves damage grading {weber}.",
    "**Human accuracy assumed perfect.** Human-reviewed cases are counted as correct.",
])

# VII. Conclusion
heading("Conclusion and Future Work")
para("This paper presented ResQRoute, a learned controller that decides, per building, whether to accept a lightweight "
     "damage prediction, escalate to a larger model, defer, or request human review, using uncertainty, image quality, "
     "device resources and asymmetric safety costs. On xBD it matched the macro-F1 of always running a large model at "
     "one-fifth to one-ninth of the CPU time, stayed within the device budget in 99.5% of decisions, and reduced "
     "dangerous severe-to-no-damage errors compared with a fixed-confidence cascade. Future work will add pre/post "
     "fusion to the base models, train the controller on out-of-fold predictions, evaluate on real device and network "
     "traces and on unseen disaster events, and add an out-of-distribution detector so that unfamiliar regions are "
     "routed to human review rather than trusted on confidence alone {hendrycks}.")

p = para("Data availability: xBD is available under CC BY-NC-SA 4.0. Code, crop manifests, model weights and logs: "
         "https://github.com/SriharshithaAnkareddy25/ResQRoute-", italic=True, size=8, indent=False, before=4)

# References
heading("References")
for i, k in enumerate(_order, 1):
    p = para(indent=False, after=1)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    add_runs(p, f"[{i}]\t{REFS[k]}", size=8)

unused = set(REFS) - set(_order)
assert not unused, f"uncited references: {unused}"
doc.save(OUT)
print(f"wrote {OUT} ({len(_order)} references)")
